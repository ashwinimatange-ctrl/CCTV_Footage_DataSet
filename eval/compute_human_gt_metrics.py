#!/usr/bin/env python3
"""Compute MAE/RMSE/MAPE from filled human manual_count vs auto_count.

Fails if manual_count cells are empty — does not invent gold GT.

Env:
  HUMAN_GT_DIR — folder with aggregate_blank_filled.csv (or aggregate_blank.csv)
  OUT_DIR — default artifacts/eval
"""
from __future__ import annotations

import csv
import json
import math
import os
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
OUT = Path(os.environ.get("OUT_DIR", str(REPO / "artifacts" / "eval")))
SRC_DIR = Path(os.environ.get("HUMAN_GT_DIR", str(REPO / "data" / "human_gt")))


def load_rows() -> tuple[list[dict], Path]:
    candidates = [
        SRC_DIR / "aggregate_blank_filled.csv",
        SRC_DIR / "aggregate_blank.csv",
    ]
    for p in candidates:
        if p.is_file() and p.stat().st_size > 0:
            rows = list(csv.DictReader(p.open(encoding="utf-8")))
            if rows:
                return rows, p
    xlsx = SRC_DIR / "Human_Gold_GT_Workbook.xlsx"
    if xlsx.is_file():
        from openpyxl import load_workbook

        wb = load_workbook(xlsx, data_only=True)
        if "Aggregate_blank" not in wb.sheetnames:
            raise SystemExit("Workbook missing Aggregate_blank sheet")
        ws = wb["Aggregate_blank"]
        header_row = None
        for r in range(1, 5):
            vals = [ws.cell(r, c).value for c in range(1, 20)]
            if vals and "manual_count" in [str(v) if v is not None else "" for v in vals]:
                header_row = r
                break
        if header_row is None:
            raise SystemExit("Could not find manual_count header in Aggregate_blank")
        headers = []
        c = 1
        while True:
            v = ws.cell(header_row, c).value
            if v is None and c > 1:
                break
            headers.append(str(v))
            c += 1
        rows = []
        r = header_row + 1
        while True:
            first = ws.cell(r, 1).value
            if first is None and ws.cell(r, 2).value is None:
                break
            row = {headers[i]: ws.cell(r, i + 1).value for i in range(len(headers))}
            rows.append({k: ("" if v is None else v) for k, v in row.items()})
            r += 1
        return rows, xlsx
    raise SystemExit(
        f"No aggregate data under {SRC_DIR}. Provide aggregate_blank_filled.csv "
        "or set HUMAN_GT_DIR."
    )


def to_float(v) -> float | None:
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def mae(xs: list[float]) -> float:
    return sum(abs(x) for x in xs) / len(xs) if xs else float("nan")


def rmse(xs: list[float]) -> float:
    return math.sqrt(sum(x * x for x in xs) / len(xs)) if xs else float("nan")


def mape(pairs: list[tuple[float, float]]) -> float:
    vals = [abs(a - m) / abs(m) for a, m in pairs if m != 0]
    return 100.0 * sum(vals) / len(vals) if vals else float("nan")


def main() -> int:
    rows, src = load_rows()
    missing = []
    pairs_all: list[tuple[float, float]] = []
    by_class: dict[str, list[tuple[float, float]]] = defaultdict(list)
    by_dir: dict[str, list[tuple[float, float]]] = defaultdict(list)

    for i, r in enumerate(rows, 1):
        m = to_float(r.get("manual_count"))
        a = to_float(r.get("auto_count"))
        if m is None:
            missing.append(i)
            continue
        if a is None:
            raise SystemExit(f"Row {i}: auto_count missing")
        pairs_all.append((a, m))
        by_class[str(r.get("class", ""))].append((a, m))
        by_dir[str(r.get("direction", ""))].append((a, m))

    if missing:
        raise SystemExit(
            f"REFUSING metrics: {len(missing)} row(s) still have empty manual_count. "
            "Complete human annotation first."
        )

    def pack(pairs: list[tuple[float, float]]) -> dict:
        errs = [a - m for a, m in pairs]
        return {
            "n": len(pairs),
            "MAE": round(mae(errs), 4),
            "RMSE": round(rmse(errs), 4),
            "MAPE_pct": round(mape(pairs), 4),
        }

    metrics = {
        "generated": datetime.now(timezone.utc).isoformat(),
        "source": str(src),
        "label": "human_gold_gt",
        "overall": pack(pairs_all),
        "by_class": {k: pack(v) for k, v in sorted(by_class.items())},
        "by_direction": {k: pack(v) for k, v in sorted(by_dir.items())},
    }

    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "human_gt_metrics.json").write_text(json.dumps(metrics, indent=2), encoding="utf-8")
    md = [
        "# Human gold GT metrics",
        "",
        f"Generated: {metrics['generated']}",
        f"Source: `{src}`",
        "",
        f"- n = {metrics['overall']['n']}",
        f"- MAE = {metrics['overall']['MAE']}",
        f"- RMSE = {metrics['overall']['RMSE']}",
        f"- MAPE (%) = {metrics['overall']['MAPE_pct']}",
        "",
    ]
    (OUT / "HUMAN_GT_METRICS.md").write_text("\n".join(md), encoding="utf-8")
    print(json.dumps(metrics["overall"], indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
